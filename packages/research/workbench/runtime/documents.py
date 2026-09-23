"""Structured extraction and PDF page rendering for the managed research Python."""
import argparse
import csv
import json
from pathlib import Path


def extract(path):
    suffix = path.suffix.lower()
    if suffix == '.pdf':
        from pypdf import PdfReader
        return [{'text': page.extract_text() or '', 'locator': {'page': index + 1}} for index, page in enumerate(PdfReader(path).pages)]
    if suffix == '.docx':
        from docx import Document
        document = Document(path)
        chunks = [{'text': paragraph.text, 'locator': {'paragraph': index + 1}} for index, paragraph in enumerate(document.paragraphs) if paragraph.text.strip()]
        for index, table in enumerate(document.tables):
            chunks.append({'text': '\n'.join('\t'.join(cell.text for cell in row.cells) for row in table.rows), 'locator': {'key': 'table:' + str(index + 1)}})
        return chunks
    if suffix == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as stream:
            return [{'text': json.dumps(row, ensure_ascii=False), 'locator': {'line': index + 1}} for index, row in enumerate(csv.reader(stream))]
    if suffix == '.json':
        with path.open(encoding='utf-8-sig') as stream:
            value = json.load(stream)
        rows = value.items() if isinstance(value, dict) else enumerate(value) if isinstance(value, list) else [('$', value)]
        return [{'text': json.dumps(value, ensure_ascii=False), 'locator': {'key': str(key)}} for key, value in rows]
    lines = path.read_text(encoding='utf-8-sig').splitlines()
    return [{'text': '\n'.join(lines[index:index + 40]), 'locator': {'line': index + 1}} for index in range(0, len(lines), 40)]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('action', choices=['extract', 'render'])
    parser.add_argument('path')
    parser.add_argument('--output')
    parser.add_argument('--max-pages', type=int, default=0,
                        help='Cap rendered pages for visual review; 0 renders every page')
    args = parser.parse_args()
    path = Path(args.path)
    if args.action == 'extract':
        print(json.dumps(extract(path), ensure_ascii=False))
    else:
        import pypdfium2 as pdfium
        destination = Path(args.output)
        destination.mkdir(parents=True, exist_ok=True)
        document = pdfium.PdfDocument(str(path))
        limit = len(document) if args.max_pages <= 0 else min(len(document), args.max_pages)
        pages = []
        for index in range(limit):
            target = destination / ('page-%03d.png' % (index + 1))
            document[index].render(scale=1.5).to_pil().save(target)
            pages.append(str(target))
        print(json.dumps(pages))


if __name__ == '__main__':
    main()
